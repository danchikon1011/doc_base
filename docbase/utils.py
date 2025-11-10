from __future__ import annotations

import html
import mimetypes
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
from xml.etree import ElementTree as ET
from zipfile import ZipFile

MSO_THEME_COLOR_INDEX = None

DOCX_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
DOCX_W = f"{{{DOCX_W_NS}}}"
DOCX_NS_MAP = {"w": DOCX_W_NS}

DOCX_NUMFMT_STYLE_MAP: Dict[str, Dict[str, str]] = {
    "decimal": {"css": "decimal", "type": "1"},
    "decimalzero": {"css": "decimal-leading-zero"},
    "decimalleadingzero": {"css": "decimal-leading-zero"},
    "upperroman": {"css": "upper-roman", "type": "I"},
    "lowerroman": {"css": "lower-roman", "type": "i"},
    "upperletter": {"css": "upper-alpha", "type": "A"},
    "lowerletter": {"css": "lower-alpha", "type": "a"},
    "ordinal": {"css": "decimal"},
    "ordtext": {"css": "decimal"},
    "cardinaltext": {"css": "decimal"},
    "hex": {"css": "decimal"},
    "decimalfullwidth": {"css": "decimal"},
    "decimalhalfwidth": {"css": "decimal"},
}

DOCX_BULLET_CHAR_STYLES: Dict[str, str] = {
    "•": "disc",
    "●": "disc",
    "▪": "square",
    "■": "square",
    "○": "circle",
    "◦": "circle",
    "–": "disc",
    "-": "disc",
    "‣": "disc",
    "»": "disc",
}

try:  # pragma: no cover - optional dependency
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover - gracefully handled later
    DocxDocument = None
    DocxParagraph = None
    WD_ALIGN_PARAGRAPH = None
    WD_COLOR_INDEX = None
    WD_UNDERLINE = None
else:  # pragma: no cover - optional dependency
    from docx.enum.text import (
        WD_ALIGN_PARAGRAPH,
        WD_COLOR_INDEX,
        WD_UNDERLINE,
    )
    from docx.text.paragraph import Paragraph as DocxParagraph

    try:  # pragma: no cover - optional dependency
        from docx.enum.dml import MSO_THEME_COLOR_INDEX as _MSO_THEME_COLOR_INDEX
    except Exception:  # pragma: no cover - gracefully handled later
        _MSO_THEME_COLOR_INDEX = None
    MSO_THEME_COLOR_INDEX = _MSO_THEME_COLOR_INDEX

try:  # pragma: no cover - optional dependency
    from pptx import Presentation
except Exception:  # pragma: no cover
    Presentation = None

try:  # pragma: no cover - optional dependency
    from openpyxl import Workbook, load_workbook
except Exception:  # pragma: no cover
    Workbook = None
    load_workbook = None

try:  # pragma: no cover - optional dependency
    import textract
except Exception:  # pragma: no cover
    textract = None

try:  # pragma: no cover - optional dependency
    from pdfminer.high_level import extract_text as extract_pdf_text
except Exception:  # pragma: no cover
    extract_pdf_text = None

ALLOWED_EXTENSIONS = {
    "txt",
    "md",
    "markdown",
    "doc",
    "docx",
    "pdf",
    "ppt",
    "pptx",
    "xls",
    "xlsx",
    "csv",
    "rtf",
}

EDITABLE_EXTENSIONS = ALLOWED_EXTENSIONS - {"pdf"}


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _extract_with_textract(path: Path) -> Optional[str]:
    if textract is None:
        return None
    try:
        return textract.process(str(path)).decode("utf-8", errors="ignore")
    except Exception:
        return None


def extract_text_from_file(path: Path, extension: str) -> str:
    extension = extension.lower()
    if extension in {"txt", "md", "markdown", "csv"}:
        return read_text_file(path)
    if extension in {"docx"}:
        paragraphs = extract_docx_paragraphs(path)
        if paragraphs:
            return "\n".join(paragraphs)
    if extension == "doc":
        extracted = _extract_with_textract(path)
        if extracted is not None:
            return extracted
    if extension in {"ppt", "pptx"}:
        slides = extract_pptx_slides(path)
        if slides:
            return "\n\n".join(slides)
    if extension in {"xls", "xlsx"}:
        sheets = extract_spreadsheet_sheets(path)
        if sheets:
            flattened: List[str] = []
            for name, rows in sheets:
                flattened.append(f"[{name}]")
                for row in rows:
                    flattened.append("\t".join(cell for cell in row))
            return "\n".join(flattened)
    if extension == "rtf":
        extracted = _extract_with_textract(path)
        if extracted is not None:
            return extracted
        return read_text_file(path)
    if extension == "pdf":
        if extract_pdf_text is not None:
            try:
                text = extract_pdf_text(str(path))
                if text and text.strip():
                    return text
            except Exception:
                pass
        extracted = _extract_with_textract(path)
        if extracted is not None:
            return extracted
        try:
            return path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return ""
    extracted = _extract_with_textract(path)
    if extracted is not None:
        return extracted
    return read_text_file(path)


def ensure_directory(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def write_content_to_file(content: str, extension: str, destination: Path) -> Optional[Path]:
    extension = extension.lower()
    if extension in {"txt", "md", "markdown", "csv"}:
        destination.write_text(content, encoding="utf-8")
        return destination
    if extension in {"doc", "docx"} and DocxDocument is not None:
        doc = DocxDocument()
        for paragraph in content.splitlines() or [""]:
            doc.add_paragraph(paragraph)
        doc.save(destination)
        return destination
    if extension in {"ppt", "pptx"} and Presentation is not None:
        presentation = Presentation()
        layout = presentation.slide_layouts[1]
        slide = presentation.slides.add_slide(layout)
        if slide.shapes.title:
            slide.shapes.title.text = "Документ"
        placeholder = None
        if len(slide.shapes.placeholders) > 1:
            placeholder = slide.shapes.placeholders[1]
        if placeholder and hasattr(placeholder, "text"):
            placeholder.text = content
        else:
            textbox = slide.shapes.add_textbox(left=0, top=0, width=presentation.slide_width, height=presentation.slide_height)
            textbox.text_frame.text = content
        presentation.save(destination)
        return destination
    if extension in {"xls", "xlsx"} and Workbook is not None:
        workbook = Workbook()
        sheet = workbook.active
        for idx, line in enumerate(content.splitlines(), start=1):
            sheet.cell(row=idx, column=1, value=line)
        workbook.save(destination)
        return destination
    if extension == "rtf":
        destination.write_text(content, encoding="utf-8")
        return destination
    destination.write_text(content, encoding="utf-8")
    return destination


def chunk_text(text: str, max_length: int = 600) -> Iterable[str]:
    words = text.split()
    chunk = []
    count = 0
    for word in words:
        chunk.append(word)
        count += len(word) + 1
        if count >= max_length:
            yield " ".join(chunk)
            chunk = []
            count = 0
    if chunk:
        yield " ".join(chunk)


def detect_mime_type(filename: str) -> str:
    mime_type, _ = mimetypes.guess_type(filename)
    return mime_type or "application/octet-stream"


DOCX_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PPTX_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
SPREADSHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


DOCX_THEME_COLOR_MAP = {
    "ACCENT_1": "#4472C4",
    "ACCENT_2": "#ED7D31",
    "ACCENT_3": "#A5A5A5",
    "ACCENT_4": "#FFC000",
    "ACCENT_5": "#5B9BD5",
    "ACCENT_6": "#70AD47",
    "BACKGROUND_1": "#FFFFFF",
    "BACKGROUND_2": "#E7E6E6",
    "DARK_1": "#000000",
    "DARK_2": "#44546A",
    "FOLLOWED_HYPERLINK": "#954F72",
    "HYPERLINK": "#0563C1",
    "LIGHT_1": "#FFFFFF",
    "LIGHT_2": "#E7E6E6",
    "TEXT_1": "#000000",
    "TEXT_2": "#44546A",
}

DOCX_HIGHLIGHT_COLOR_MAP = {
    "BLACK": "#000000",
    "BLUE": "#0000FF",
    "BRIGHT_GREEN": "#00FF00",
    "DARK_BLUE": "#00008B",
    "DARK_RED": "#8B0000",
    "DARK_YELLOW": "#808000",
    "GRAY_25": "#C0C0C0",
    "GRAY_50": "#7F7F7F",
    "GREEN": "#008000",
    "PINK": "#FFC0CB",
    "RED": "#FF0000",
    "TURQUOISE": "#40E0D0",
    "VIOLET": "#EE82EE",
    "WHITE": "#FFFFFF",
    "YELLOW": "#FFFF00",
}


def _zip_xml(path: Path, member: str) -> Optional[ET.Element]:
    try:
        with ZipFile(path) as archive:
            with archive.open(member) as payload:
                return ET.fromstring(payload.read())
    except KeyError:
        return None
    except Exception:
        return None


def _first_defined(*values: Any) -> Any:
    for value in values:
        if value is not None:
            return value
    return None


def _docx_format_hex(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    stripped = value.strip().lstrip("#")
    if not stripped:
        return None
    if len(stripped) not in {3, 6}:
        return None
    if any(ch not in "0123456789abcdefABCDEF" for ch in stripped):
        return None
    return f"#{stripped.upper()}"


def _docx_color_to_css(color: Any) -> Optional[str]:
    if color is None:
        return None
    rgb = getattr(color, "rgb", None)
    if rgb is not None:
        try:
            return _docx_format_hex(str(rgb))
        except Exception:
            pass
    theme_color = getattr(color, "theme_color", None)
    if theme_color is not None:
        name = getattr(theme_color, "name", None) or str(theme_color)
        if name:
            name = name.split()[0].upper()
            css = DOCX_THEME_COLOR_MAP.get(name)
            if css:
                return css
    if MSO_THEME_COLOR_INDEX is not None and rgb is None:
        try:
            theme_value = getattr(color, "theme_color", None)
            if theme_value is not None:
                resolved = MSO_THEME_COLOR_INDEX(theme_value)
                name = getattr(resolved, "name", None)
                if name and name in DOCX_THEME_COLOR_MAP:
                    return DOCX_THEME_COLOR_MAP[name]
        except Exception:
            pass
    return None


def _docx_highlight_to_css(value: Any) -> Optional[str]:
    if value is None:
        return None
    name = getattr(value, "name", None)
    if name is None:
        if isinstance(value, str):
            name = value.upper()
        elif WD_COLOR_INDEX is not None:
            try:
                resolved = WD_COLOR_INDEX(value)
                name = getattr(resolved, "name", None)
            except Exception:
                name = None
    if not name:
        return None
    name = name.split()[0].upper()
    return DOCX_HIGHLIGHT_COLOR_MAP.get(name)


def _docx_length_to_css(length: Any) -> Optional[str]:
    if length is None:
        return None
    value: Optional[float] = None
    if hasattr(length, "pt"):
        try:
            value = float(length.pt)
        except Exception:
            value = None
    if value is None:
        try:
            value = float(length)
        except Exception:
            value = None
    if value is None or value <= 0:
        return None
    formatted = f"{value:.2f}".rstrip("0").rstrip(".")
    return f"{formatted}pt"


def _docx_collect_font_candidates(run: Any, defaults: Optional[dict]) -> List[Any]:
    fonts: List[Any] = []
    font = getattr(run, "font", None)
    if font is not None:
        fonts.append(font)
    style = getattr(run, "style", None)
    style_font = getattr(style, "font", None) if style is not None else None
    if style_font is not None:
        fonts.append(style_font)
    paragraph = getattr(run, "paragraph", None)
    if paragraph is not None:
        para_style = getattr(paragraph, "style", None)
        para_font = getattr(para_style, "font", None) if para_style is not None else None
        if para_font is not None:
            fonts.append(para_font)
    if defaults:
        for fallback_font in defaults.get("font_chain", []) or []:
            if fallback_font is not None:
                fonts.append(fallback_font)
    return fonts


def _docx_font_property(fonts: Sequence[Any], attribute: str) -> Any:
    for font in fonts:
        if font is None:
            continue
        try:
            value = getattr(font, attribute)
        except Exception:
            continue
        if value is not None:
            return value
    return None


def _docx_is_underlined(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return value.lower() != "none"
    if value in (False, 0):
        return False
    if WD_UNDERLINE is not None:
        try:
            if value == WD_UNDERLINE.NONE:
                return False
        except Exception:
            pass
    if hasattr(value, "value") and WD_UNDERLINE is not None:
        try:
            if value.value == WD_UNDERLINE.NONE:
                return False
        except Exception:
            pass
    return True


def _docx_document_defaults(document: Any) -> dict:
    defaults = {
        "font_chain": [],
        "font_name": None,
        "font_size": None,
        "font_color": None,
        "font_highlight": None,
    }
    styles = getattr(document, "styles", None)
    if styles is not None:
        normal_style = None
        try:
            normal_style = styles["Normal"]
        except Exception:
            normal_style = None
        if normal_style is not None:
            normal_font = getattr(normal_style, "font", None)
            if normal_font is not None:
                defaults["font_chain"].append(normal_font)
                defaults["font_name"] = getattr(normal_font, "name", None)
                defaults["font_size"] = getattr(normal_font, "size", None)
                defaults["font_color"] = getattr(normal_font, "color", None)
                defaults["font_highlight"] = getattr(
                    normal_font, "highlight_color", None
                )
    return defaults


def _docx_apply_defaults(html_text: str, defaults: Optional[dict]) -> str:
    if not defaults:
        return html_text
    style_parts: List[str] = []
    font_name = defaults.get("font_name")
    font_size = defaults.get("font_size")
    font_color = defaults.get("font_color")
    highlight = defaults.get("font_highlight")
    if font_name:
        safe_name = font_name.replace("\\", "\\\\").replace('"', '\\"')
        style_parts.append(f'font-family: "{safe_name}"')
    size_css = _docx_length_to_css(font_size)
    if size_css:
        style_parts.append(f"font-size: {size_css}")
    color_css = _docx_color_to_css(font_color)
    if color_css:
        style_parts.append(f"color: {color_css}")
    highlight_css = _docx_highlight_to_css(highlight)
    if highlight_css:
        style_parts.append(f"background-color: {highlight_css}")
    if not style_parts:
        return html_text
    style_attr = html.escape("; ".join(style_parts), quote=True)
    return f"<span style=\"{style_attr}\">{html_text}</span>"


def _docx_extract_bullet_style(text: Optional[str]) -> Optional[str]:
    if not text:
        return None
    stripped = text.replace("%1", "").strip()
    if not stripped:
        return None
    return DOCX_BULLET_CHAR_STYLES.get(stripped)


def _docx_numbering_level_info(element: Any) -> Dict[str, Any]:
    info: Dict[str, Any] = {}
    if element is None:
        return info
    num_fmt_elem = element.find("w:numFmt", DOCX_NS_MAP)
    lvl_text_elem = element.find("w:lvlText", DOCX_NS_MAP)
    start_elem = element.find("w:start", DOCX_NS_MAP)
    fmt_val = None
    if num_fmt_elem is not None:
        fmt_val = num_fmt_elem.get(f"{DOCX_W}val")
    if lvl_text_elem is not None:
        info["lvl_text"] = lvl_text_elem.get(f"{DOCX_W}val")
    if start_elem is not None:
        start_val = start_elem.get(f"{DOCX_W}val")
        if start_val is not None:
            try:
                info["start"] = int(start_val)
            except (TypeError, ValueError):
                pass
    fmt_lower = (fmt_val or "").strip().lower()
    if fmt_lower:
        info["num_fmt"] = fmt_lower
    lvl_text = info.get("lvl_text")
    ordered = True
    if fmt_lower in {"bullet", "none"}:
        ordered = False
    elif not fmt_lower:
        ordered = "%" in (lvl_text or "%1")
    info["ordered"] = ordered
    if ordered:
        style_info = DOCX_NUMFMT_STYLE_MAP.get(fmt_lower)
        if style_info:
            css_value = style_info.get("css")
            type_value = style_info.get("type")
            if css_value:
                info["css"] = css_value
            if type_value:
                info["type"] = type_value
    else:
        bullet_style = _docx_extract_bullet_style(lvl_text)
        if bullet_style:
            info["css"] = bullet_style
        else:
            info.setdefault("css", "disc")
    return info


def _docx_build_numbering_map(path: Path) -> Dict[int, Dict[int, Dict[str, Any]]]:
    mapping: Dict[int, Dict[int, Dict[str, Any]]] = {}
    try:
        with ZipFile(path) as archive:
            data = archive.read("word/numbering.xml")
    except (KeyError, FileNotFoundError):
        return mapping
    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        return mapping
    abstract_defs: Dict[str, Dict[int, Dict[str, Any]]] = {}
    for abstract in root.findall("w:abstractNum", DOCX_NS_MAP):
        abs_id = abstract.get(f"{DOCX_W}abstractNumId")
        if abs_id is None:
            continue
        levels: Dict[int, Dict[str, Any]] = {}
        for lvl in abstract.findall("w:lvl", DOCX_NS_MAP):
            ilvl = lvl.get(f"{DOCX_W}ilvl")
            if ilvl is None:
                continue
            try:
                level_index = int(ilvl)
            except (TypeError, ValueError):
                continue
            level_info = _docx_numbering_level_info(lvl)
            if level_info:
                levels[level_index] = level_info
        if levels:
            abstract_defs[abs_id] = levels
    for num in root.findall("w:num", DOCX_NS_MAP):
        num_id = num.get(f"{DOCX_W}numId")
        if num_id is None:
            continue
        abstract_ref_elem = num.find("w:abstractNumId", DOCX_NS_MAP)
        if abstract_ref_elem is None:
            continue
        abstract_ref = abstract_ref_elem.get(f"{DOCX_W}val")
        if abstract_ref is None:
            continue
        base_levels = abstract_defs.get(abstract_ref, {})
        levels: Dict[int, Dict[str, Any]] = {
            level: dict(info) for level, info in base_levels.items()
        }
        for override in num.findall("w:lvlOverride", DOCX_NS_MAP):
            ilvl = override.get(f"{DOCX_W}ilvl")
            if ilvl is None:
                continue
            try:
                level_index = int(ilvl)
            except (TypeError, ValueError):
                continue
            level_info = dict(levels.get(level_index, {}))
            start_override = override.find("w:startOverride", DOCX_NS_MAP)
            if start_override is not None:
                start_val = start_override.get(f"{DOCX_W}val")
                if start_val is not None:
                    try:
                        level_info["start"] = int(start_val)
                    except (TypeError, ValueError):
                        pass
            lvl = override.find("w:lvl", DOCX_NS_MAP)
            if lvl is not None:
                merged = _docx_numbering_level_info(lvl)
                if merged:
                    level_info.update(merged)
            if level_info:
                levels[level_index] = level_info
        try:
            num_id_int = int(num_id)
        except (TypeError, ValueError):
            continue
        if levels:
            mapping[num_id_int] = levels
    return mapping


def _docx_numbering_level(
    numbering_map: Dict[int, Dict[int, Dict[str, Any]]],
    num_id: int,
    level: int,
) -> Dict[str, Any]:
    base = numbering_map.get(num_id, {}).get(level, {})
    info: Dict[str, Any] = dict(base)
    info["num_id"] = num_id
    info["level"] = level
    num_fmt = info.get("num_fmt")
    if isinstance(num_fmt, str):
        info["num_fmt"] = num_fmt.lower()
    ordered = info.get("ordered")
    if ordered is None:
        fmt = info.get("num_fmt")
        ordered = fmt not in (None, "bullet", "none")
        info["ordered"] = ordered
    if info.get("css") is None:
        if info.get("ordered"):
            fmt = info.get("num_fmt")
            style_info = DOCX_NUMFMT_STYLE_MAP.get(fmt or "") if fmt else None
            if style_info and style_info.get("css"):
                info["css"] = style_info["css"]
            else:
                info["css"] = "decimal"
        else:
            info["css"] = "disc"
    if info.get("ordered") and info.get("type") is None:
        fmt = info.get("num_fmt")
        style_info = DOCX_NUMFMT_STYLE_MAP.get(fmt or "") if fmt else None
        if style_info and style_info.get("type"):
            info["type"] = style_info["type"]
    if info.get("start") is not None:
        try:
            info["start"] = int(info["start"])
        except (TypeError, ValueError):
            info["start"] = None
    return info


def _docx_list_info(paragraph, numbering_map) -> Optional[Dict[str, Any]]:
    if paragraph is None or getattr(paragraph, "_p", None) is None:
        return None
    p = paragraph._p
    pPr = getattr(p, "pPr", None)
    if pPr is None:
        return None
    numPr = getattr(pPr, "numPr", None)
    if numPr is None:
        return None
    num_id_obj = getattr(numPr, "numId", None)
    if num_id_obj is None or getattr(num_id_obj, "val", None) is None:
        return None
    try:
        num_id_val = int(num_id_obj.val)
    except (TypeError, ValueError):
        return None
    ilvl_obj = getattr(numPr, "ilvl", None)
    level_value = 0
    if ilvl_obj is not None and getattr(ilvl_obj, "val", None) is not None:
        try:
            level_value = int(ilvl_obj.val)
        except (TypeError, ValueError):
            level_value = 0
    info = _docx_numbering_level(numbering_map, num_id_val, level_value)
    lvl_text = info.get("lvl_text")
    if lvl_text is None:
        base = numbering_map.get(num_id_val, {}).get(level_value, {})
        if base.get("lvl_text") is not None:
            info["lvl_text"] = base["lvl_text"]
    return info


def _docx_trim_list_stack(
    html_parts: List[str], stack: List[Dict[str, Any]], target_length: int
) -> None:
    while len(stack) > target_length:
        current = stack.pop()
        if current.get("open_item"):
            html_parts.append("</li>")
            current["open_item"] = False
        html_parts.append(f"</{current['tag']}>")


def _docx_close_all_lists(html_parts: List[str], stack: List[Dict[str, Any]]) -> None:
    _docx_trim_list_stack(html_parts, stack, 0)


def _docx_open_list(
    html_parts: List[str], info: Dict[str, Any], depth: int
) -> Dict[str, Any]:
    ordered = bool(info.get("ordered"))
    tag = "ol" if ordered else "ul"
    classes = ["docx-list", f"docx-list-level-{depth}"]
    classes.append("docx-list--ordered" if ordered else "docx-list--bullet")
    class_attr = " ".join(classes)
    attr_parts = [f'class="{class_attr}"']
    css_value = info.get("css")
    style_attr = None
    if css_value:
        style_attr = html.escape("; ".join([f"list-style-type: {css_value}"]), quote=True)
        attr_parts.append(f'style="{style_attr}"')
    type_value = info.get("type")
    if ordered and type_value:
        attr_parts.append(f'type="{html.escape(str(type_value), quote=True)}"')
    start_value = info.get("start")
    if ordered and isinstance(start_value, int) and start_value not in (0, 1):
        attr_parts.append(f'start="{start_value}"')
    attrs = " " + " ".join(attr_parts) if attr_parts else ""
    html_parts.append(f"<{tag}{attrs}>")
    entry = {
        "tag": tag,
        "num_id": info.get("num_id"),
        "level": depth,
        "ordered": ordered,
        "css": css_value,
        "type": type_value,
        "open_item": False,
    }
    return entry


def _docx_ensure_list_context(
    html_parts: List[str],
    stack: List[Dict[str, Any]],
    info: Dict[str, Any],
    numbering_map: Dict[int, Dict[int, Dict[str, Any]]],
) -> Dict[str, Any]:
    target_length = info.get("level", 0) + 1
    _docx_trim_list_stack(html_parts, stack, target_length)
    while len(stack) < target_length:
        depth = len(stack)
        level_info = _docx_numbering_level(numbering_map, info.get("num_id", 0), depth)
        if depth == info.get("level"):
            level_info.update({k: v for k, v in info.items() if v is not None})
        if stack:
            parent = stack[-1]
            if not parent.get("open_item"):
                parent_classes = [
                    "docx-list-item",
                    f"docx-list-item--level-{parent['level']}",
                ]
                parent_classes.append(
                    "docx-list-item--ordered" if parent.get("ordered") else "docx-list-item--bullet"
                )
                html_parts.append(
                    f"<li class=\"{' '.join(parent_classes)}\">&nbsp;"
                )
                parent["open_item"] = True
        entry = _docx_open_list(html_parts, level_info, depth)
        stack.append(entry)
    current = stack[-1]
    if (
        current.get("num_id") != info.get("num_id")
        or current.get("ordered") != info.get("ordered")
        or current.get("css") != info.get("css")
        or current.get("type") != info.get("type")
    ):
        if current.get("open_item"):
            html_parts.append("</li>")
            current["open_item"] = False
        html_parts.append(f"</{current['tag']}>")
        stack.pop()
        entry = _docx_open_list(html_parts, info, info.get("level", 0))
        stack.append(entry)
        current = entry
    return current


def _docx_start_list_item(
    html_parts: List[str], current: Dict[str, Any], info: Dict[str, Any], content: str
) -> None:
    if current.get("open_item"):
        html_parts.append("</li>")
        current["open_item"] = False
    text = content or "&nbsp;"
    classes = ["docx-list-item", f"docx-list-item--level-{info.get('level', 0)}"]
    classes.append(
        "docx-list-item--ordered" if info.get("ordered") else "docx-list-item--bullet"
    )
    html_parts.append(f"<li class=\"{' '.join(classes)}\">{text}")
    current["open_item"] = True


def extract_docx_paragraphs(path: Path) -> List[str]:
    paragraphs: List[str] = []
    if DocxDocument is not None:
        try:
            document = DocxDocument(str(path))
            paragraphs = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text and paragraph.text.strip()
            ]
        except Exception:
            paragraphs = []
    if paragraphs:
        return paragraphs
    root = _zip_xml(path, "word/document.xml")
    if root is None:
        return []
    extracted: List[str] = []
    for paragraph in root.iter(f"{DOCX_NS}p"):
        texts = [node.text or "" for node in paragraph.iter(f"{DOCX_NS}t")]
        combined = "".join(texts).strip()
        if combined:
            extracted.append(combined)
    return extracted


def extract_pptx_slides(path: Path) -> List[str]:
    slides: List[str] = []
    if Presentation is not None:
        try:
            presentation = Presentation(str(path))
            for slide in presentation.slides:
                texts: List[str] = []
                for shape in slide.shapes:
                    text = getattr(shape, "text", "")
                    if text:
                        snippet = str(text).strip()
                        if snippet:
                            texts.append(snippet)
                combined = "\n".join(texts).strip()
                if combined:
                    slides.append(combined)
        except Exception:
            slides = []
    if slides:
        return slides
    try:
        with ZipFile(path) as archive:
            members = [
                name
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ]
            members.sort(
                key=lambda name: int(name.split("slide")[1].split(".xml")[0])
            )
            extracted: List[str] = []
            for member in members:
                with archive.open(member) as payload:
                    root = ET.fromstring(payload.read())
                texts = [node.text or "" for node in root.iter(f"{PPTX_NS}t")]
                combined = "\n".join(part.strip() for part in texts if (part or "").strip())
                if combined:
                    extracted.append(combined)
            return extracted
    except Exception:
        return []
    return []


def extract_spreadsheet_sheets(path: Path) -> List[Tuple[str, List[List[str]]]]:
    if load_workbook is not None:
        try:
            workbook = load_workbook(filename=str(path), data_only=True)
            sheets: List[Tuple[str, List[List[str]]]] = []
            for sheet in workbook.worksheets:
                rows: List[List[str]] = []
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if cell is None else str(cell) for cell in row]
                    rows.append(cells)
                sheets.append((sheet.title, rows))
            return sheets
        except Exception:
            pass
    try:
        with ZipFile(path) as archive:
            shared_strings: List[str] = []
            if "xl/sharedStrings.xml" in archive.namelist():
                with archive.open("xl/sharedStrings.xml") as payload:
                    root = ET.fromstring(payload.read())
                for entry in root.iter(f"{SPREADSHEET_NS}si"):
                    parts = [node.text or "" for node in entry.iter(f"{SPREADSHEET_NS}t")]
                    shared_strings.append("".join(parts))

            workbook_root = _zip_xml(path, "xl/workbook.xml")
            if workbook_root is None:
                return []

            relationships: dict[str, str] = {}
            rel_root = _zip_xml(path, "xl/_rels/workbook.xml.rels")
            if rel_root is not None:
                for rel in rel_root.iter(f"{REL_NS}Relationship"):
                    target = rel.attrib.get("Target")
                    if target:
                        relationships[rel.attrib.get("Id")] = target

            sheets: List[Tuple[str, List[List[str]]]] = []
            for sheet in workbook_root.iter(f"{SPREADSHEET_NS}sheet"):
                name = sheet.attrib.get("name", "Лист")
                rel_id = sheet.attrib.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                )
                if not rel_id:
                    continue
                target = relationships.get(rel_id)
                if not target:
                    continue
                member = target if target.startswith("xl/") else f"xl/{target}"
                if member not in archive.namelist():
                    continue
                with archive.open(member) as payload:
                    root = ET.fromstring(payload.read())
                rows: List[List[str]] = []
                for row in root.iter(f"{SPREADSHEET_NS}row"):
                    cells: List[str] = []
                    for cell in row.iter(f"{SPREADSHEET_NS}c"):
                        value = ""
                        cell_type = cell.attrib.get("t")
                        v = cell.find(f"{SPREADSHEET_NS}v")
                        if v is not None and v.text is not None:
                            if cell_type == "s":
                                try:
                                    idx = int(v.text)
                                    value = shared_strings[idx]
                                except (ValueError, IndexError):
                                    value = v.text
                            else:
                                value = v.text
                        elif cell_type == "inlineStr":
                            t = cell.find(f"{SPREADSHEET_NS}t")
                            if t is not None and t.text:
                                value = t.text
                        cells.append(value)
                    rows.append(cells)
                sheets.append((name, rows))
            return sheets
    except Exception:
        return []
    return []


def prepare_preview(extension: Optional[str], file_path: Optional[str], content: str) -> dict:
    ext = (extension or "").lower()
    path = Path(file_path) if file_path else None

    if ext == "pdf" and path and path.exists():
        return {"kind": "pdf"}

    if ext in {"doc", "docx", "rtf", "txt", "md", "markdown"}:
        if path and path.exists() and ext == "docx":
            html_preview = render_docx_html(path)
            if html_preview:
                return {"kind": "docx", "html": html_preview}
        paragraphs: List[str] = []
        if path and path.exists() and ext in {"doc", "docx"}:
            paragraphs = extract_docx_paragraphs(path)
        if not paragraphs and content:
            paragraphs = [line for line in content.splitlines() if line.strip()]
        return {"kind": "paragraphs", "paragraphs": paragraphs or [content]}

    if ext in {"ppt", "pptx"}:
        slides = []
        if path and path.exists():
            slides = build_pptx_preview(path)
        if not slides:
            slide_texts: List[str] = []
            if path and path.exists():
                slide_texts = extract_pptx_slides(path)
            if not slide_texts and content:
                slide_texts = [block.strip() for block in content.split("\n\n") if block.strip()]
            if slide_texts:
                slides = build_pptx_preview_from_text(slide_texts)
        return {"kind": "slides", "slides": slides}

    if ext in {"xls", "xlsx", "csv"}:
        sheets: List[Tuple[str, List[List[str]]]] = []
        if path and path.exists() and ext in {"xls", "xlsx"}:
            sheets = extract_spreadsheet_sheets(path)
        if not sheets and content:
            rows = [line.split("\t") for line in content.splitlines() if line]
            if rows:
                sheets = [("Лист 1", rows)]
        return {
            "kind": "sheets",
            "sheets": [{"name": name, "rows": rows} for name, rows in sheets],
        }

    text = content or ""
    return {"kind": "text", "text": text}


def _docx_run_to_html(run, defaults: Optional[dict] = None) -> str:
    text = html.escape(getattr(run, "text", "") or "")
    if text == "":
        return ""
    text = text.replace("\n", "<br />")

    fonts = _docx_collect_font_candidates(run, defaults)

    bold = _first_defined(getattr(run, "bold", None), _docx_font_property(fonts, "bold"))
    italic = _first_defined(
        getattr(run, "italic", None), _docx_font_property(fonts, "italic")
    )
    underline_value = _first_defined(
        getattr(run, "underline", None), _docx_font_property(fonts, "underline")
    )
    strike_value = _first_defined(
        getattr(run, "strike", None), _docx_font_property(fonts, "strike")
    )
    double_strike = _docx_font_property(fonts, "double_strike")

    style_parts: List[str] = []
    text_decorations: List[str] = []

    font_name = _docx_font_property(fonts, "name")
    if font_name is None and defaults:
        font_name = defaults.get("font_name")
    if font_name:
        safe_name = font_name.replace("\\", "\\\\").replace('"', '\\"')
        style_parts.append(f'font-family: "{safe_name}"')

    font_size_obj = _docx_font_property(fonts, "size")
    if font_size_obj is None and defaults:
        font_size_obj = defaults.get("font_size")
    size_css = _docx_length_to_css(font_size_obj)
    if size_css:
        style_parts.append(f"font-size: {size_css}")

    color_obj = _docx_font_property(fonts, "color")
    if color_obj is None and defaults:
        color_obj = defaults.get("font_color")
    color_css = _docx_color_to_css(color_obj)
    if color_css:
        style_parts.append(f"color: {color_css}")

    highlight_obj = _docx_font_property(fonts, "highlight_color")
    if highlight_obj is None and defaults:
        highlight_obj = defaults.get("font_highlight")
    highlight_css = _docx_highlight_to_css(highlight_obj)
    if highlight_css:
        style_parts.append(f"background-color: {highlight_css}")

    all_caps = _docx_font_property(fonts, "all_caps")
    if all_caps:
        style_parts.append("text-transform: uppercase")

    small_caps = _docx_font_property(fonts, "small_caps")
    if small_caps:
        style_parts.append("font-variant: small-caps")

    shadow = _docx_font_property(fonts, "shadow")
    if shadow:
        style_parts.append("text-shadow: 0.5px 0.5px 0 currentColor")

    if strike_value:
        text_decorations.append("line-through")
    if double_strike:
        text_decorations.append("line-through")

    underline = _docx_is_underlined(underline_value)
    if underline:
        text_decorations.append("underline")

    if text_decorations:
        unique: List[str] = []
        for decoration in text_decorations:
            if decoration not in unique:
                unique.append(decoration)
        style_parts.append(f"text-decoration: {' '.join(unique)}")

    superscript = _docx_font_property(fonts, "superscript")
    subscript = _docx_font_property(fonts, "subscript")
    if superscript and not subscript:
        style_parts.append("vertical-align: super")
        if size_css is None:
            style_parts.append("font-size: 0.83em")
    elif subscript and not superscript:
        style_parts.append("vertical-align: sub")
        if size_css is None:
            style_parts.append("font-size: 0.83em")

    style_attr = ""
    if style_parts:
        style_attr = ' style="{}"'.format(html.escape("; ".join(style_parts), quote=True))

    class_attr = ""
    if underline:
        class_attr = ' class="docx-underline"'

    if style_attr or class_attr:
        text = f"<span{class_attr}{style_attr}>{text}</span>"

    if bold:
        text = f"<strong>{text}</strong>"
    if italic:
        text = f"<em>{text}</em>"
    return text
def _docx_alignment_class(paragraph) -> str:
    if WD_ALIGN_PARAGRAPH is None:
        return ""
    alignment = getattr(paragraph, "alignment", None)
    if alignment is None:
        return ""
    if alignment == WD_ALIGN_PARAGRAPH.CENTER:
        return " docx-align-center"
    if alignment == WD_ALIGN_PARAGRAPH.RIGHT:
        return " docx-align-right"
    if alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
        return " docx-align-justify"
    return ""


def _docx_heading_level(style_name: str) -> int:
    match = re.search(r"(\d+)", style_name)
    if match:
        try:
            value = int(match.group(1))
            return max(1, min(6, value))
        except ValueError:
            return 2
    return 2


def _docx_render_table(table, defaults: Optional[dict]) -> str:
    rows_html: List[str] = []
    for row in getattr(table, "rows", []):
        cells_html: List[str] = []
        for cell in getattr(row, "cells", []):
            cell_paragraphs: List[str] = []
            for paragraph in getattr(cell, "paragraphs", []):
                paragraph_html = "".join(
                    _docx_run_to_html(run, defaults)
                    for run in getattr(paragraph, "runs", [])
                )
                if not paragraph_html:
                    raw_text = html.escape(getattr(paragraph, "text", "") or "")
                    raw_text = raw_text.replace("\n", "<br />")
                    paragraph_html = _docx_apply_defaults(raw_text, defaults)
                else:
                    paragraph_html = paragraph_html.replace("\n", "<br />")
                if paragraph_html:
                    cell_paragraphs.append(paragraph_html)
            cells_html.append("<td>{}</td>".format("<br />".join(cell_paragraphs) or "&nbsp;"))
        if cells_html:
            rows_html.append(f"<tr>{''.join(cells_html)}</tr>")
    if not rows_html:
        return ""
    return f"<table class=\"docx-table\"><tbody>{''.join(rows_html)}</tbody></table>"


def render_docx_html(path: Path) -> str:
    if DocxDocument is None or DocxParagraph is None:
        return ""
    try:
        document = DocxDocument(str(path))
    except Exception:
        return ""

    html_parts: List[str] = []
    list_stack: List[Dict[str, Any]] = []
    defaults = _docx_document_defaults(document)
    numbering_map = _docx_build_numbering_map(path)

    paragraphs: Sequence = getattr(document, "paragraphs", [])
    tables: Sequence = getattr(document, "tables", [])
    tables_iter = iter(tables)
    table_map = {getattr(table, "_tbl", None): table for table in tables}

    body = getattr(document, "element", None)
    body = getattr(body, "body", None)
    if body is None:
        return ""

    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "tbl":
            table = table_map.get(child)
            if table is None:
                try:
                    table = next(tables_iter)
                except StopIteration:
                    table = None
            if table is not None:
                _docx_close_all_lists(html_parts, list_stack)
                table_html = _docx_render_table(table, defaults)
                if table_html:
                    html_parts.append(table_html)
            continue
        if tag != "p":
            continue
        paragraph = DocxParagraph(child, document)
        paragraph_html = "".join(
            _docx_run_to_html(run, defaults) for run in paragraph.runs
        )
        if not paragraph_html:
            raw_text = html.escape(paragraph.text or "")
            raw_text = raw_text.replace("\n", "<br />")
            paragraph_html = _docx_apply_defaults(raw_text, defaults)
        paragraph_html = paragraph_html.replace("\n", "<br />")
        paragraph_text = paragraph_html.strip()
        style = getattr(paragraph, "style", None)
        style_name = (getattr(style, "name", "") or "").lower()
        if style_name.startswith("heading"):
            if not paragraph_text:
                continue
            _docx_close_all_lists(html_parts, list_stack)
            level = _docx_heading_level(style_name)
            html_parts.append(f"<h{level}>{paragraph_text}</h{level}>")
            continue
        list_info = _docx_list_info(paragraph, numbering_map)
        if list_info is not None:
            current_list = _docx_ensure_list_context(
                html_parts, list_stack, list_info, numbering_map
            )
            _docx_start_list_item(
                html_parts,
                current_list,
                list_info,
                paragraph_html if paragraph_text else "",
            )
            continue
        if not paragraph_text:
            continue
        _docx_close_all_lists(html_parts, list_stack)
        alignment_class = _docx_alignment_class(paragraph)
        html_parts.append(
            f"<p class=\"docx-paragraph{alignment_class}\">{paragraph_text}</p>"
        )

    _docx_close_all_lists(html_parts, list_stack)
    return "".join(html_parts)


def _split_slide_lines(lines: Sequence[str]) -> Tuple[Optional[str], List[str]]:
    cleaned = [line.strip() for line in lines if line.strip()]
    if not cleaned:
        return None, []
    if len(cleaned) == 1:
        return cleaned[0], []
    return cleaned[0], cleaned[1:]


def build_pptx_preview_from_text(slides: Sequence[str]) -> List[dict]:
    rendered: List[dict] = []
    for slide in slides:
        lines = slide.splitlines()
        title, bullet_lines = _split_slide_lines(lines)
        if not bullet_lines and not title:
            continue
        body: List[str] = []
        if bullet_lines:
            body.append("<ul class=\"pptx-list\">")
            body.extend(f"<li>{html.escape(line)}</li>" for line in bullet_lines)
            body.append("</ul>")
        rendered.append({
            "title": title,
            "html": "".join(body) if body else None,
        })
    return rendered


def build_pptx_preview(path: Path) -> List[dict]:
    if Presentation is None:
        return []
    try:
        presentation = Presentation(str(path))
    except Exception:
        return []
    slides_payload: List[dict] = []
    for slide in presentation.slides:
        title_shape = getattr(slide.shapes, "title", None)
        title = ""
        if title_shape is not None:
            title = (getattr(title_shape, "text", "") or "").strip()
        body_parts: List[str] = []
        list_stack: List[int] = []

        def close_lists(level: int = 0) -> None:
            while len(list_stack) > level:
                body_parts.append("</ul>")
                list_stack.pop()

        for shape in slide.shapes:
            if shape == title_shape:
                continue
            text_frame = getattr(shape, "text_frame", None)
            if text_frame is None:
                continue
            for paragraph in text_frame.paragraphs:
                runs = getattr(paragraph, "runs", [])
                paragraph_html = "".join(_docx_run_to_html(run) for run in runs)
                if not paragraph_html:
                    paragraph_html = html.escape(getattr(paragraph, "text", "") or "")
                paragraph_html = paragraph_html.replace("\n", "<br />").strip()
                if not paragraph_html:
                    continue
                level = getattr(paragraph, "level", 0) or 0
                bullet = getattr(paragraph, "_p", None)
                bullet_enabled = False
                if bullet is not None:
                    pPr = getattr(bullet, "pPr", None)
                    if pPr is not None:
                        bullet_enabled = any(
                            child.tag.split("}")[-1].startswith("bu") and child.tag.split("}")[-1] != "buNone"
                            for child in pPr.iterchildren()
                        )
                if bullet_enabled or level > 0:
                    while len(list_stack) <= level:
                        body_parts.append("<ul class=\"pptx-list\">")
                        list_stack.append(len(list_stack))
                    close_lists(level + 1)
                    body_parts.append(f"<li>{paragraph_html}</li>")
                else:
                    close_lists(0)
                    body_parts.append(f"<p>{paragraph_html}</p>")
        close_lists(0)
        content_html = "".join(body_parts)
        slides_payload.append(
            {
                "title": title or None,
                "html": content_html or None,
            }
        )
    return slides_payload
