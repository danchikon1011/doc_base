"""Helpers for saving files and extracting textual content."""
from __future__ import annotations

import shutil
from pathlib import Path
from typing import Tuple

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from PyPDF2 import PdfReader

from ..core.config import get_settings

settings = get_settings()


SUPPORTED_EXTENSIONS = {".docx", ".pptx", ".xlsx", ".pdf"}


def save_upload_file(uploaded_file, destination: Path) -> Path:
    """Persist the uploaded file to the storage folder."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("wb") as buffer:
        shutil.copyfileobj(uploaded_file.file, buffer)
    return destination


def extract_text_from_file(path: Path) -> str:
    """Extract raw text from the supported document types."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        doc = DocxDocument(path)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)
    if suffix == ".pptx":
        presentation = Presentation(path)
        texts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
        return "\n".join(texts)
    if suffix == ".xlsx":
        wb = load_workbook(filename=path, read_only=True)
        cells = []
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                row_values = [str(cell) for cell in row if cell is not None]
                if row_values:
                    cells.append("\t".join(row_values))
        return "\n".join(cells)
    if suffix == ".pdf":
        reader = PdfReader(path)
        texts = []
        for page in reader.pages:
            texts.append(page.extract_text() or "")
        return "\n".join(texts)
    return ""


def build_storage_path(filename: str) -> Path:
    return settings.file_storage_path / filename
